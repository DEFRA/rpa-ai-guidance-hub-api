"""Reading the box Word draws by bordering the paragraphs themselves.

The third form a box arrives in, and the only one with no container: the author
borders a few adjacent body paragraphs and the box is that they share it. These cases
are about assembling that run, about what does and does not belong to it, and about
it landing where a text box in the same place lands.

All fixture text is invented, as everywhere in this package.
"""

from collections.abc import Callable
from typing import Any

from docx.document import Document
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.text.paragraph import Paragraph

from app.guidance.parsing import parser

HEADING = "Applying"

# The `docx_bytes` fixture's contract, restated here as it is in test_textboxes.py.
Build = Callable[[Document], None]
DocxBytes = Callable[[Build], bytes]

_SIDES = ("w:top", "w:left", "w:bottom", "w:right")


def _bordered_content(docx_bytes: DocxBytes, build: Build) -> str:
    """The Markdown of the section `build` writes its paragraphs into."""

    def with_heading(document: Document) -> None:
        document.add_heading(HEADING, level=1)
        build(document)

    return parser.parse_docx(docx_bytes(with_heading)).sections[0].content


def _border(
    paragraph: Paragraph,
    *sides: str,
    value: str = "single",
    colour: str | None = None,
) -> Paragraph:
    """Draw a border on the sides named, or on all four when none are."""
    border = OxmlElement("w:pBdr")
    for side in sides or _SIDES:
        edge = OxmlElement(side)
        edge.set(qn("w:val"), value)
        if colour is not None:
            edge.set(qn("w:color"), colour)
        border.append(edge)

    paragraph._p.get_or_add_pPr().append(border)
    return paragraph


def _boxed(document: Document, *texts: str, style: str | None = None) -> None:
    """Add one bordered paragraph per text given, which is one box between them."""
    for text in texts:
        _border(document.add_paragraph(text, style=style))


class TestAssemblingTheBox:
    def test_bordered_paragraphs_become_a_blockquote(self, docx_bytes):
        """A border on all four sides is the box a reader sees, and Word draws the
        same box with a one-cell table and with the drawing tools."""

        def build(document):
            _boxed(document, "Claim reference: xxxxx")

        assert _bordered_content(docx_bytes, build) == "> Claim reference: xxxxx"

    def test_adjacent_bordered_paragraphs_are_one_box(self, docx_bytes):
        """The box is the run, not the paragraph: Word writes the border on each
        paragraph it encloses and draws one rectangle round the lot."""

        def build(document):
            _boxed(document, "Claim reference: xxxxx", "Task type: signoff")

        assert _bordered_content(docx_bytes, build) == (
            "> Claim reference: xxxxx\n>\n> Task type: signoff"
        )

    def test_the_first_unbordered_paragraph_ends_it(self, docx_bytes):
        """Nothing else says where the box stops, and the guidance resuming under a
        template reads as the last line of the template until it does."""

        def build(document):
            document.add_paragraph("Add a note to the case.")
            _boxed(document, "Claim reference: xxxxx")
            document.add_paragraph("Tell the processor the outcome.")

        assert _bordered_content(docx_bytes, build) == (
            "Add a note to the case.\n\n"
            "> Claim reference: xxxxx\n\n"
            "Tell the processor the outcome."
        )

    def test_two_boxes_apart_are_two_boxes(self, docx_bytes):
        """Two runs separated by prose are two quotes, not one - the prose between
        them is what the second box is for."""

        def build(document):
            _boxed(document, "Pass - no action needed.")
            document.add_paragraph("Where rework is needed instead:")
            _boxed(document, "Rework required.")

        assert _bordered_content(docx_bytes, build) == (
            "> Pass - no action needed.\n\n"
            "Where rework is needed instead:\n\n"
            "> Rework required."
        )

    def test_a_blank_bordered_paragraph_is_spacing_and_not_a_line(self, docx_bytes):
        """Word spaces the inside of a box with empty paragraphs, which carry the
        border like the rest. Quoting them would put bare markers in the block."""

        def build(document):
            _boxed(document, "Claim reference: xxxxx", "", "Task type: signoff")

        assert _bordered_content(docx_bytes, build) == (
            "> Claim reference: xxxxx\n>\n> Task type: signoff"
        )

    def test_a_list_inside_a_box_is_still_a_list(self, docx_bytes):
        """A box holds blocks, not text, so its paragraphs go through the same
        machinery as the body's - reading it as text drops every marker."""

        def build(document):
            _boxed(document, "Include:", style="List Bullet")

        assert _bordered_content(docx_bytes, build) == "> - Include:"


class TestWhereOneBoxEndsAndTheNextBegins:
    def test_a_change_of_border_starts_another_box(self, docx_bytes):
        """Word grows one frame down consecutive paragraphs only while the border
        stays the same, and draws a fresh one the moment it changes. It is how these
        guides put an email's subject line in a box above the box holding its body,
        and two boxes on the page are two blockquotes."""

        def build(document):
            _border(document.add_paragraph("A quality check was carried out on"))
            _border(document.add_paragraph("SBI: 123456789"), value="double")

        assert _bordered_content(docx_bytes, build) == (
            "> A quality check was carried out on\n\n> SBI: 123456789"
        )

    def test_a_border_that_renders_the_same_still_starts_another(self, docx_bytes):
        """`auto` is black and so is an explicit black, and Word draws them apart
        anyway - which is the whole of what separates the two boxes in every guide
        that draws a subject line above a body."""

        def build(document):
            _border(document.add_paragraph("HOLD806 - rework required"), colour="auto")
            _border(document.add_paragraph("Claim reference: xxxxx"), colour="000000")

        assert _bordered_content(docx_bytes, build) == (
            "> HOLD806 - rework required\n\n> Claim reference: xxxxx"
        )

    def test_an_unchanged_border_is_still_one_box(self, docx_bytes):
        """The paragraphs of one box carry the same border, which is what makes them
        one box. Splitting on every paragraph would put a break between every line
        of a case note."""

        def build(document):
            _boxed(document, "Claim reference: xxxxx", "Task type: signoff")

        assert _bordered_content(docx_bytes, build) == (
            "> Claim reference: xxxxx\n>\n> Task type: signoff"
        )


class TestWhatIsNotABox:
    def test_a_border_on_one_side_is_a_rule(self, docx_bytes):
        """Word's own Title style underlines itself with a bottom border. Quoting a
        heading because of it would be a loss dressed up as a gain."""

        def build(document):
            _border(document.add_paragraph("Section summary"), "w:bottom")

        assert _bordered_content(docx_bytes, build) == "Section summary"

    def test_a_side_turned_off_is_not_a_line(self, docx_bytes):
        """Word writes the side and says none rather than dropping the element, so
        four sides present is not four sides drawn."""

        def build(document):
            _border(document.add_paragraph("Section summary"), value="none")

        assert _bordered_content(docx_bytes, build) == "Section summary"

    def test_a_bordered_heading_still_opens_its_section(self, docx_bytes):
        """A border round a heading is emphasis. Taking it into a box would lose the
        section, and everything under it with the section."""

        def build(document):
            _border(document.add_heading("Rework required", level=2))
            document.add_paragraph("Attach the completed checklist.")

        sections = parser.parse_docx(docx_bytes(build)).sections
        assert [section.heading for section in sections] == ["Rework required"]
        assert sections[0].content == "Attach the completed checklist."


class TestWhereTheBoxLands:
    def test_a_box_between_two_items_belongs_to_the_item_above_it(self, docx_bytes):
        """The same rule a text box drawn there gets, and for the same reason:
        closing the run reopens the items after it at the margin."""

        def build(document):
            document.add_paragraph("Send the form", style="List Bullet")
            _boxed(document, "Check the version")
            document.add_paragraph("Keep a copy", style="List Bullet")

        assert _bordered_content(docx_bytes, build) == (
            "- Send the form\n\n  > Check the version\n- Keep a copy"
        )

    def test_a_box_closing_the_document_is_still_filed(self, docx_bytes):
        """The run is held until something ends it, and the end of the body is one
        of the things that can."""

        def build(document):
            _boxed(document, "No further action required.")

        assert _bordered_content(docx_bytes, build) == "> No further action required."

    def test_a_table_after_a_box_does_not_swallow_it(self, docx_bytes: Any):
        """A table closes the run of bordered paragraphs before it is rendered, so
        the box stands ahead of it rather than after it."""

        def build(document):
            _boxed(document, "Claim reference: xxxxx")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Option"
            table.cell(0, 1).text = "Area"

        assert _bordered_content(docx_bytes, build) == (
            "> Claim reference: xxxxx\n\n| Option | Area |\n| --- | --- |"
        )
