"""The offline parse script, as far as the ids it lays a document out under.

All fixture text is invented, as everywhere in these tests.
"""

from __future__ import annotations

import sys
import uuid

import docx
import parse_docx


def _guide(path):
    document = docx.Document()
    document.add_heading("Applying", level=1)
    document.add_paragraph("Fill in the form.")
    document.save(path)


class TestTheIdsItMints:
    def test_a_new_document_and_its_version_are_laid_out_under_time_ordered_uuids(
        self, tmp_path, monkeypatch
    ):
        """As the service mints them, so a layout made offline is one the service
        could have made."""
        source = tmp_path / "Applying.docx"
        _guide(source)
        output = tmp_path / "bucket"
        monkeypatch.setattr(
            sys, "argv", ["parse_docx.py", str(source), "--output-dir", str(output)]
        )

        assert parse_docx.main() == 0

        (document_dir,) = output.iterdir()
        version_dirs = [
            path for path in document_dir.iterdir() if path.name != "assets"
        ]
        assert uuid.UUID(document_dir.name).version == 7
        assert [uuid.UUID(path.name).version for path in version_dirs] == [7]
